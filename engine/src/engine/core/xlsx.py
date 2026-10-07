from collections.abc import Sequence
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from engine.core.errors import EngineError, OutputWriteError
from engine.core.messages import t
from engine.core.paths import extended_path, write_failure_text

FORMULA_PREFIX = "="
MAX_COLUMN_WIDTH = 60
MIN_COLUMN_WIDTH = 8
COLUMN_PADDING = 2
QUANTITY_DECIMALS = 2

type CellValue = str | int | float


class TemplateError(EngineError):
    pass


def open_sheet(template: Path | None, sheet_name: str | None) -> tuple[Workbook, Worksheet]:
    """Opens the user's template, or a blank workbook when none is given."""
    if template is None:
        workbook = Workbook()
        return workbook, _active_sheet(workbook)
    workbook = _load_template(template)
    if not sheet_name:
        return workbook, _active_sheet(workbook)
    if sheet_name not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        raise TemplateError(
            t("xlsx.sheet_missing", sheet=sheet_name),
            file=template,
            hint=t("xlsx.sheet_missing_hint", available=available),
        )
    return workbook, workbook[sheet_name]


def write_row(
    sheet: Worksheet, row: int, values: Sequence[CellValue], *, bold: bool = False
) -> None:
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        _keep_text_literal(cell, value)
        if bold:
            cell.font = Font(bold=True)


def write_table(
    sheet: Worksheet,
    headers: Sequence[CellValue],
    rows: Sequence[Sequence[CellValue]],
    *,
    header_row: int = 1,
    fit: bool = True,
) -> None:
    """Bold header then one row per line; `fit` is off when a template owns the layout."""
    write_row(sheet, header_row, headers, bold=True)
    for offset, row in enumerate(rows, start=1):
        write_row(sheet, header_row + offset, row)
    if fit:
        sheet.freeze_panes = f"A{header_row + 1}"
        fit_columns(sheet, [headers, *rows])


def quantity_cell(value: float) -> CellValue:
    return int(value) if value.is_integer() else value


def quantity_text(value: float) -> str:
    return str(int(value)) if value.is_integer() else str(round(value, QUANTITY_DECIMALS))


def fit_columns(sheet: Worksheet, rows: Sequence[Sequence[CellValue]]) -> None:
    for column in range(1, max((len(row) for row in rows), default=0) + 1):
        longest = max((len(str(row[column - 1])) for row in rows if len(row) >= column), default=0)
        width = min(max(longest + COLUMN_PADDING, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH)
        sheet.column_dimensions[get_column_letter(column)].width = width


def save_workbook(workbook: Workbook, target: Path) -> None:
    try:
        extended_path(target.parent).mkdir(parents=True, exist_ok=True)
        workbook.save(extended_path(target))
    except OSError as error:
        message, hint = write_failure_text(
            target, error, t("xlsx.save_failed"), t("xlsx.save_failed_hint")
        )
        raise OutputWriteError(message, file=target, hint=hint) from error


def _keep_text_literal(cell: Cell, value: CellValue) -> None:
    # Attribute text starting with "=" must never be evaluated as a formula.
    if isinstance(value, str) and value.startswith(FORMULA_PREFIX):
        cell.data_type = "s"


def _load_template(template: Path) -> Workbook:
    try:
        return load_workbook(template)
    except (OSError, BadZipFile, InvalidFileException, KeyError) as error:
        raise TemplateError(
            t("xlsx.template_unreadable"), file=template, hint=t("xlsx.template_unreadable_hint")
        ) from error


def _active_sheet(workbook: Workbook) -> Worksheet:
    sheet = workbook.active
    if not isinstance(sheet, Worksheet):
        raise TemplateError(t("xlsx.no_sheet"))
    return sheet
