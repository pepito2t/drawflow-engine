"""A quick look at a file handed to the assistant: what it is, and which feature fits it."""

from collections import Counter
from contextlib import closing
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from engine.assistant.messages import t
from engine.core.cache import FileCache, resolve_cache_root
from engine.core.errors import InvalidInputError
from engine.core.settings import load_general_settings

MAX_FILE_BYTES = 200_000_000
MAX_PDF_LINES = 40
MAX_TEXT_CHARS = 2_000
MAX_SHEET_ROWS = 5
MAX_SHEET_CELLS = 12
MAX_SHEETS = 5
MAX_PARAGRAPHS = 15
MAX_BLOCK_NAMES = 15
PLAN_SUFFIXES = {".dwg", ".dxf"}
SUGGESTED_FEATURE = {
    ".dwg": "dwg-parts",
    ".dxf": "dwg-parts",
    ".pdf": "pdf-report",
    ".xlsx": "soumission",
}


def inspect_file(settings: Path, raw_path: str) -> dict[str, Any]:
    path = Path(raw_path)
    if not path.is_absolute() or not path.is_file():
        raise InvalidInputError(
            t("inspect.file_not_found"), file=path, hint=t("inspect.file_not_found.hint")
        )
    if path.stat().st_size > MAX_FILE_BYTES:
        raise InvalidInputError(t("inspect.file_too_large"), file=path)
    suffix = path.suffix.lower()
    base = {
        "file": str(path),
        "name": path.name,
        "type": suffix.lstrip("."),
        "size_bytes": path.stat().st_size,
        "suggested_feature": SUGGESTED_FEATURE.get(suffix),
    }
    if suffix in PLAN_SUFFIXES:
        return {**base, **_plan(path, settings)}
    if suffix == ".pdf":
        return {**base, **_pdf(path)}
    if suffix == ".xlsx":
        return {**base, **_workbook(path)}
    if suffix == ".docx":
        return {**base, **_document(path)}
    raise InvalidInputError(
        t("inspect.unsupported_type"), file=path, hint=t("inspect.unsupported_type.hint")
    )


def _plan(path: Path, settings: Path) -> dict[str, Any]:
    from engine.parts.oda import DXF_SUFFIX, is_dxf
    from engine.parts.worker import DXF_CACHE_NAMESPACE, extract_file

    cache_root = resolve_cache_root(load_general_settings(settings))
    # A DWG conversion takes longer than the assistant waits for a tool: only a plan the feature
    # already converted is read here.
    if not is_dxf(path) and not FileCache(cache_root, DXF_CACHE_NAMESPACE).has(path, DXF_SUFFIX):
        return {"plan_read": False, "note": t("inspect.plan_not_read")}
    extraction = extract_file(path, oda_executable=None, cache_root=cache_root)
    blocks = Counter(part.block for part in extraction.parts)
    attributes = sorted({tag for part in extraction.parts for tag in part.attributes})
    return {
        "plan_read": True,
        "blocks_total": sum(blocks.values()),
        "blocks": [
            {"name": name, "count": count} for name, count in blocks.most_common(MAX_BLOCK_NAMES)
        ],
        "attributes": attributes,
        "warnings": [warning.message for warning in extraction.warnings],
    }


def _pdf(path: Path) -> dict[str, Any]:
    from engine.modules.pdf_report.reader import count_pages, read_pages

    pages = count_pages(path)
    with closing(read_pages(path)) as reader:
        first_page = next(reader, None)
    first_lines = [line.text for line in first_page.lines[:MAX_PDF_LINES]] if first_page else []
    text = "\n".join(first_lines)[:MAX_TEXT_CHARS]
    return {"pages": pages, "first_page_text": text, "scanned": pages > 0 and not text}


def _workbook(path: Path) -> dict[str, Any]:
    from engine.modules.soumission.reader import open_workbook, sheet_rows

    workbook = open_workbook(path, str(path))
    try:
        sheets = []
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            rows = [
                [str(cell) for cell in row[:MAX_SHEET_CELLS] if cell is not None and str(cell)]
                for row in sheet_rows(sheet, str(path), MAX_SHEET_ROWS)
            ]
            sheets.append({"sheet": sheet.title, "first_rows": [row for row in rows if row]})
    finally:
        workbook.close()
    return {"sheets": sheets}


def _document(path: Path) -> dict[str, Any]:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(str(path))
    except (PackageNotFoundError, BadZipFile, KeyError, OSError) as error:
        raise InvalidInputError(
            t("inspect.unreadable_document"), file=path, hint=t("inspect.unreadable_document.hint")
        ) from error
    paragraphs = [
        paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()
    ]
    return {
        "paragraphs": paragraphs[:MAX_PARAGRAPHS],
        "tables": len(document.tables),
    }
