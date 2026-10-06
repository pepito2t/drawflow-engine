from collections.abc import Iterator, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import IO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook

from engine.core.anomalies import Anomaly
from engine.core.errors import EngineError
from engine.core.fields import KeyValue
from engine.core.text import fold
from engine.modules.soumission.numbers import parse_amount
from engine.modules.soumission.settings import (
    MIN_HEADER_MATCHES,
    SoumissionSettings,
    synonyms_of,
)

HEADER_PUNCTUATION = " .:"

type CellValue = str | float
type Row = dict[str, CellValue]


class WorkbookReadError(EngineError):
    pass


@dataclass(frozen=True)
class SubmissionTable:
    source: str
    sheet: str
    rows: list[Row]


@dataclass(frozen=True)
class WorkbookContent:
    tables: list[SubmissionTable] = field(default_factory=list)
    warnings: list[Anomaly] = field(default_factory=list)


def read_workbook(
    stream: Path | IO[bytes], source: str, settings: SoumissionSettings
) -> WorkbookContent:
    workbook = _open(stream, source)
    try:
        tables: list[SubmissionTable] = []
        warnings: list[Anomaly] = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            table, sheet_warnings = _read_sheet(rows, sheet.title, source, settings)
            warnings.extend(sheet_warnings)
            if table is not None:
                tables.append(table)
    finally:
        workbook.close()
    if not tables:
        warnings.append(
            Anomaly(
                "Aucun tableau de soumission reconnu (en-têtes introuvables).",
                hint="Ajoutez les en-têtes de ce fichier aux synonymes de colonnes dans "
                "Paramètres → Soumission.",
            )
        )
    return WorkbookContent(tables=tables, warnings=warnings)


def _open(stream: Path | IO[bytes], source: str) -> Workbook:
    try:
        return load_workbook(stream, read_only=True, data_only=True)
    except (OSError, BadZipFile, InvalidFileException, KeyError) as error:
        raise WorkbookReadError(
            "Le classeur Excel est illisible.",
            file=Path(source),
            hint="Vérifiez qu'il s'ouvre dans Excel (format .xlsx).",
        ) from error


def _read_sheet(
    rows: Sequence[Sequence[object]], title: str, source: str, settings: SoumissionSettings
) -> tuple[SubmissionTable | None, list[Anomaly]]:
    header = _find_header(rows, settings)
    if header is None:
        return None, []
    header_index, positions = header
    numeric = settings.numeric_column_names()
    parsed: list[Row] = []
    unreadable = 0
    for raw in rows[header_index + 1 :]:
        row, row_unreadable = _normalize_row(raw, positions, numeric)
        unreadable += row_unreadable
        if row:
            parsed.append(row)
    table = SubmissionTable(source, title, parsed)
    if not unreadable:
        return table, []
    return table, [
        Anomaly(
            f"{unreadable} montant(s) illisible(s) gardé(s) en texte.",
            location=f"feuille « {title} »",
            hint="Vérifiez les colonnes converties en nombres dans Paramètres → Soumission.",
        )
    ]


def _find_header(
    rows: Sequence[Sequence[object]], settings: SoumissionSettings
) -> tuple[int, dict[str, int]] | None:
    for index, row in enumerate(rows[: settings.header_search_rows]):
        positions = _match_columns(row, settings.columns)
        if len(positions) >= MIN_HEADER_MATCHES:
            return index, positions
    return None


def _match_columns(row: Sequence[object], columns: Sequence[KeyValue]) -> dict[str, int]:
    cells = {
        fold(str(value)).strip(HEADER_PUNCTUATION): position for position, value in _texts(row)
    }
    positions: dict[str, int] = {}
    for column in columns:
        for synonym in synonyms_of(column):
            position = cells.get(fold(synonym).strip(HEADER_PUNCTUATION))
            if position is not None:
                positions[column.key] = position
                break
    return positions


def _texts(row: Sequence[object]) -> Iterator[tuple[int, str]]:
    for position, value in enumerate(row):
        if isinstance(value, str) and value.strip():
            yield position, value


def _normalize_row(
    raw: Sequence[object], positions: dict[str, int], numeric: set[str]
) -> tuple[Row, int]:
    row: Row = {}
    unreadable = 0
    for column, position in positions.items():
        value = raw[position] if position < len(raw) else None
        if value is None or (isinstance(value, str) and not value.strip()):
            continue
        if column in numeric:
            amount = parse_amount(value)
            if amount is None:
                unreadable += 1
                row[column] = str(value).strip()
                continue
            row[column] = amount
        else:
            row[column] = str(value).strip()
    return row, unreadable
