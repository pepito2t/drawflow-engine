from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from engine.modules.soumission.export import export_submissions
from engine.modules.soumission.reader import SubmissionTable

COLUMNS = ["Position", "Désignation", "Total"]


def test_rows_from_all_tables_share_one_sheet_with_their_origin(tmp_path: Path) -> None:
    tables = [
        SubmissionTable(
            "a.xlsx", "Offre", [{"Position": "1", "Désignation": "Panneau", "Total": 10.0}]
        ),
        SubmissionTable("b.pdf > lot.xlsx", "Lot", [{"Position": "=2", "Total": 5.5}]),
    ]
    target = tmp_path / "Sortie é" / "soumission.xlsx"

    count = export_submissions(tables, COLUMNS, target)

    sheet = load_workbook(target).active
    assert isinstance(sheet, Worksheet)
    assert count == 2
    assert [list(row) for row in sheet.iter_rows(values_only=True)] == [
        ["Position", "Désignation", "Total", "Fichier source", "Feuille"],
        ["1", "Panneau", 10, "a.xlsx", "Offre"],
        ["=2", None, 5.5, "b.pdf > lot.xlsx", "Lot"],
    ]
