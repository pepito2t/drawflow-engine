from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

from engine.modules.soumission.reader import WorkbookReadError, read_workbook
from engine.modules.soumission.settings import SoumissionSettings
from engine.modules.soumission.tests.workbooks import submission_bytes, write_submission


def test_finds_the_header_and_normalizes_rows(tmp_path: Path) -> None:
    path = write_submission(tmp_path / "offre é.xlsx")

    content = read_workbook(path, str(path), SoumissionSettings())

    [table] = content.tables
    assert table.sheet == "Offre"
    assert table.rows[0] == {
        "Position": "1.1",
        "Désignation": "Panneau composite RAL 7016",
        "Quantité": 12.0,
        "Unité": "m2",
        "Prix unitaire": 1234.5,
    }
    assert table.rows[1]["Total"] == 140.0
    assert len(table.rows) == 3


def test_unreadable_amounts_stay_as_text_with_a_warning() -> None:
    content = read_workbook(BytesIO(submission_bytes()), "offre.xlsx", SoumissionSettings())

    assert content.tables[0].rows[2]["Quantité"] == "forfait"
    assert content.tables[0].rows[2]["Prix unitaire"] == 2000.0
    [warning] = content.warnings
    assert warning.message == "1 montant(s) illisible(s) gardé(s) en texte."
    assert warning.location == "feuille « Offre »"
    assert warning.hint is not None


def test_header_beyond_the_search_window_is_not_found() -> None:
    content = read_workbook(
        BytesIO(submission_bytes(title_rows=5)),
        "offre.xlsx",
        SoumissionSettings(header_search_rows=3),
    )

    assert content.tables == []
    assert "Aucun tableau" in content.warnings[-1].message


def test_date_cells_are_written_as_day_month_year() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Pos.", "Désignation", "Qté", "Unité", "P.U.", "Montant"])
    sheet.append(["1.1", datetime(2024, 1, 1, 0, 0), 1, "pce", 10, date(2024, 3, 15)])
    buffer = BytesIO()
    workbook.save(buffer)

    content = read_workbook(BytesIO(buffer.getvalue()), "offre.xlsx", SoumissionSettings())

    [row] = content.tables[0].rows
    assert row["Désignation"] == "01.01.2024"
    assert row["Total"] == "15.03.2024"
    assert [warning.message for warning in content.warnings] == [
        "1 montant(s) illisible(s) gardé(s) en texte."
    ]


def test_unreadable_workbook_is_reported(tmp_path: Path) -> None:
    path = tmp_path / "cassé.xlsx"
    path.write_text("pas un classeur", encoding="utf-8")

    with pytest.raises(WorkbookReadError) as caught:
        read_workbook(path, str(path), SoumissionSettings())

    assert caught.value.file == path
