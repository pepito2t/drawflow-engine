import json
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from engine.modules.soumission.tests.workbooks import submission_bytes, write_submission
from engine.testing.pdf import PdfSpec, TextItem, write_pdf

GOLDEN_RELATIVE_PATH = Path("fixtures") / "soumission" / "expected" / "soumission.json"


def repository_root() -> Path:
    for parent in Path(__file__).resolve().parents:
        if (parent / "fixtures").is_dir() and (parent / "engine").is_dir():
            return parent
    raise RuntimeError("repository root not found")


@pytest.fixture
def sources(tmp_path: Path) -> list[Path]:
    lot_a = tmp_path / "Lot A"
    lot_b = tmp_path / "Lot B" / "Reçus"
    lot_b.mkdir(parents=True)
    lot_a.mkdir()
    original = write_submission(lot_a / "Façadier.xlsx")
    shutil.copyfile(original, lot_b / "Façadier (copie).xlsx")
    spec = PdfSpec(
        pages=[[TextItem(50, 50, "Offre")]], attachments={"menuiserie.xlsx": submission_bytes()}
    )
    write_pdf(lot_b / "Menuisier.pdf", spec)
    return [lot_a, tmp_path / "Lot B"]


def run_soumission(tmp_path: Path, folders: list[Path]) -> list[dict[str, Any]]:
    input_file = tmp_path / "inputs.json"
    payload = {
        "folders": [str(folder) for folder in folders],
        "project": "Tour B",
        "output_folder": str(tmp_path / "Sortie"),
        "preview": False,
    }
    input_file.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    completed = subprocess.run(
        [sys.executable, "-m", "engine.cli", "run", "soumission", "--input", str(input_file)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    return [json.loads(line) for line in completed.stdout.splitlines()]


def test_submissions_from_several_folders_match_golden_file(
    tmp_path: Path, sources: list[Path]
) -> None:
    events = run_soumission(tmp_path, sources)

    result = events[-1]
    assert result["type"] == "result", events
    sheet = load_workbook(result["outputs"][0]).active
    assert isinstance(sheet, Worksheet)
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    golden = json.loads((repository_root() / GOLDEN_RELATIVE_PATH).read_text(encoding="utf-8"))
    assert rows == golden["rows"]


def test_duplicate_copies_are_reported(tmp_path: Path, sources: list[Path]) -> None:
    events = run_soumission(tmp_path, sources)

    duplicates = [
        event for event in events if event["type"] == "warning" and "Doublon" in event["message"]
    ]
    assert len(duplicates) == 1
    assert duplicates[0]["file"].endswith("Façadier (copie).xlsx")


def test_missing_file_from_a_preset_is_a_warning_naming_it(tmp_path: Path) -> None:
    present = write_submission(tmp_path / "Façadier.xlsx")
    gone = tmp_path / "Disparue.xlsx"
    input_file = tmp_path / "inputs.json"
    payload = {
        "files": [str(gone), str(present)],
        "project": "Tour B",
        "output_folder": str(tmp_path / "Sortie"),
        "preview": False,
    }
    input_file.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    completed = subprocess.run(
        [sys.executable, "-m", "engine.cli", "run", "soumission", "--input", str(input_file)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )
    events = [json.loads(line) for line in completed.stdout.splitlines()]

    assert completed.returncode == 0, completed.stderr
    [missing] = [event for event in events if event.get("file") == str(gone)]
    assert missing["type"] == "warning" and "introuvable" in missing["message"]
    assert "préréglage" in missing["hint"]
    assert events[-1]["type"] == "result"
