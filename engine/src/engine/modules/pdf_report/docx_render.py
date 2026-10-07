from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docxtpl import DocxTemplate
from jinja2 import Environment, StrictUndefined, TemplateError, UndefinedError

from engine.core.errors import EngineError, OutputWriteError
from engine.core.paths import extended_path, write_failure_text
from engine.modules.pdf_report.messages import t

DEFAULT_TEMPLATE_NAME = "rapport-par-defaut.docx"
FIX_TEMPLATE_HINT = t("docx.fix_template_hint")


class ReportTemplateError(EngineError):
    pass


def render_report(template: Path | None, context: dict[str, Any], target: Path) -> None:
    if template is not None:
        _render(template, context, target, source=template)
        return
    with TemporaryDirectory(prefix="drawflow-docx-") as workdir:
        default = write_default_template(Path(workdir) / DEFAULT_TEMPLATE_NAME)
        _render(default, context, target, source=None)


def write_default_template(path: Path) -> Path:
    document = Document()
    document.add_heading(t("docx.default.title"), level=0)
    document.add_paragraph(
        t(
            "docx.default.summary",
            project="{{ projet }}",
            date="{{ date }}",
            count="{{ nb_plans }}",
        )
    )
    document.add_paragraph("{%p for plan in plans %}")
    document.add_heading(
        t("docx.default.plan_heading", file="{{ plan.fichier }}", pages="{{ plan.pages }}"),
        level=1,
    )
    document.add_paragraph("{%p for champ in plan.liste_champs %}")
    document.add_paragraph(
        t("docx.default.field", label="{{ champ.libelle }}", value="{{ champ.valeur }}")
    )
    document.add_paragraph("{%p endfor %}")
    document.add_paragraph(t("docx.default.references", references="{{ plan.references }}"))
    document.add_paragraph("{%p endfor %}")
    document.save(str(path))
    return path


def _render(template: Path, context: dict[str, Any], target: Path, source: Path | None) -> None:
    document = _open(template)
    try:
        document.render(context, jinja_env=Environment(undefined=StrictUndefined, autoescape=True))
    except UndefinedError as error:
        raise ReportTemplateError(
            t("docx.unknown_tag", detail=error.message),
            file=source,
            hint=FIX_TEMPLATE_HINT,
        ) from error
    except TemplateError as error:
        raise ReportTemplateError(
            t("docx.malformed_tag", detail=error.message),
            file=source,
            hint=t("docx.malformed_tag.hint"),
        ) from error
    _save(document, target)


def _open(template: Path) -> DocxTemplate:
    try:
        document = DocxTemplate(str(template))
        document.init_docx()
    except (OSError, BadZipFile, KeyError, ValueError, PackageNotFoundError) as error:
        raise ReportTemplateError(
            t("docx.unreadable"), file=template, hint=t("docx.unreadable.hint")
        ) from error
    return document


def _save(document: DocxTemplate, target: Path) -> None:
    try:
        extended_path(target.parent).mkdir(parents=True, exist_ok=True)
        document.save(str(extended_path(target)))
    except OSError as error:
        message, hint = write_failure_text(
            target, error, t("docx.save_failed"), t("docx.save_failed.hint")
        )
        raise OutputWriteError(message, file=target, hint=hint) from error
