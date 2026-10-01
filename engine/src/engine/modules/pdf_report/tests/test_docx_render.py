from datetime import datetime
from pathlib import Path

import pytest
from docx import Document

from engine.modules.pdf_report.docx_render import ReportTemplateError, render_report
from engine.modules.pdf_report.report import PlanData, build_context
from engine.modules.pdf_report.settings import DEFAULT_FIELDS

MOMENT = datetime(2026, 10, 2, 8, 0)
PLANS = [
    PlanData("Nord.pdf", 2, {"projet": "Tour B", "indice": "C"}, ["EQ-40", "P-1200"]),
    PlanData("Sud.pdf", 1, {"projet": "Tour B"}, []),
]


def paragraphs(path: Path) -> list[str]:
    return [paragraph.text for paragraph in Document(str(path)).paragraphs if paragraph.text]


def custom_template(path: Path, *lines: str) -> Path:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    document.save(str(path))
    return path


def test_default_template_lists_plans_fields_and_references(tmp_path: Path) -> None:
    target = tmp_path / "Sortie é" / "rapport.docx"

    render_report(None, build_context(PLANS, DEFAULT_FIELDS, "Tour B", MOMENT), target)

    text = paragraphs(target)
    assert "Projet : Tour B — Date : 02.10.2026 — 2 plan(s)" in text
    assert "Nord.pdf (2 page(s))" in text
    assert "Indice : C" in text
    assert "Date : " in text
    assert "Références : EQ-40, P-1200" in text


def test_user_template_can_use_field_tags_directly(tmp_path: Path) -> None:
    template = custom_template(
        tmp_path / "modèle.docx",
        "{%p for plan in plans %}",
        "{{ plan.fichier }} indice {{ plan.indice }}",
        "{%p endfor %}",
    )
    target = tmp_path / "rapport.docx"
    plans = [PlanData("Nord.pdf", 1, {"indice": "C"})]

    render_report(template, build_context(plans, DEFAULT_FIELDS, "", MOMENT), target)

    assert paragraphs(target) == ["Nord.pdf indice C"]
    assert paragraphs(template)[1] == "{{ plan.fichier }} indice {{ plan.indice }}"


def test_unknown_tag_is_named(tmp_path: Path) -> None:
    template = custom_template(tmp_path / "modèle.docx", "{{ client }}")

    with pytest.raises(ReportTemplateError) as caught:
        render_report(
            template, build_context(PLANS, DEFAULT_FIELDS, "", MOMENT), tmp_path / "r.docx"
        )

    assert "client" in caught.value.message
    assert caught.value.file == template


def test_malformed_tag_is_reported(tmp_path: Path) -> None:
    template = custom_template(tmp_path / "modèle.docx", "{{ projet ")

    with pytest.raises(ReportTemplateError):
        render_report(
            template, build_context(PLANS, DEFAULT_FIELDS, "", MOMENT), tmp_path / "r.docx"
        )


def test_unreadable_template_is_reported(tmp_path: Path) -> None:
    template = tmp_path / "modèle.docx"
    template.write_text("pas un docx", encoding="utf-8")

    with pytest.raises(ReportTemplateError) as caught:
        render_report(
            template, build_context(PLANS, DEFAULT_FIELDS, "", MOMENT), tmp_path / "r.docx"
        )

    assert caught.value.file == template


def test_values_are_escaped_for_word(tmp_path: Path) -> None:
    plans = [PlanData("A&B <1>.pdf", 1, {})]
    target = tmp_path / "rapport.docx"

    render_report(None, build_context(plans, DEFAULT_FIELDS, "R&D", MOMENT), target)

    assert "A&B <1>.pdf (1 page(s))" in paragraphs(target)
