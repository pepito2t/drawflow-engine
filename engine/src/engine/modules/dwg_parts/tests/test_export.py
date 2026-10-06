from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from engine.core.xlsx import TemplateError
from engine.modules.dwg_parts.export import export_parts
from engine.parts.aggregation import PartLine
from engine.parts.settings import PartsListSettings

HEADERS = ("Bloc", "Référence")
LINES = [
    PartLine(("PANNEAU", "P-900"), 1, ("a.dwg",)),
    PartLine(("PANNEAU", "P-1200"), 2.5, ("a.dwg", "b.dwg")),
]


def sheet_values(path: Path, sheet: str | None = None) -> list[tuple[object, ...]]:
    workbook = load_workbook(path)
    worksheet = workbook[sheet] if sheet else workbook.active
    assert isinstance(worksheet, Worksheet)
    return [tuple(row) for row in worksheet.iter_rows(values_only=True)]


def make_template(path: Path) -> Path:
    workbook = Workbook()
    cover = workbook.active
    assert isinstance(cover, Worksheet)
    cover.title = "Garde"
    cover["A1"] = "Projet Tour B"
    workbook.create_sheet("Pièces")
    workbook.save(path)
    return path


def test_default_layout_writes_header_lines_and_totals(tmp_path: Path) -> None:
    target = tmp_path / "Sortie é" / "liste.xlsx"

    export_parts(LINES, HEADERS, target, None, PartsListSettings())

    assert sheet_values(target) == [
        ("Bloc", "Référence", "Quantité", "Plans"),
        ("PANNEAU", "P-900", 1, "a.dwg"),
        ("PANNEAU", "P-1200", 2.5, "a.dwg, b.dwg"),
    ]


def test_template_sheet_and_header_row_are_respected(tmp_path: Path) -> None:
    template = make_template(tmp_path / "modèle.xlsx")
    target = tmp_path / "liste.xlsx"
    settings = PartsListSettings(template_sheet="Pièces", header_row=3)

    export_parts(LINES, HEADERS, target, template, settings)

    rows = sheet_values(target, "Pièces")
    assert rows[2] == ("Bloc", "Référence", "Quantité", "Plans")
    assert rows[3][1] == "P-900"
    assert sheet_values(target, "Garde")[0] == ("Projet Tour B",)
    assert sheet_values(template, "Pièces") == []


def test_unknown_sheet_lists_available_ones(tmp_path: Path) -> None:
    template = make_template(tmp_path / "modèle.xlsx")

    with pytest.raises(TemplateError) as caught:
        export_parts(
            LINES, HEADERS, tmp_path / "x.xlsx", template, PartsListSettings(template_sheet="Nope")
        )

    assert caught.value.hint is not None
    assert "Garde, Pièces" in caught.value.hint


def test_unreadable_template_is_reported(tmp_path: Path) -> None:
    template = tmp_path / "modèle.xlsx"
    template.write_text("pas un classeur", encoding="utf-8")

    with pytest.raises(TemplateError) as caught:
        export_parts(LINES, HEADERS, tmp_path / "x.xlsx", template, PartsListSettings())

    assert caught.value.file == template


def test_text_starting_with_equals_is_never_a_formula(tmp_path: Path) -> None:
    target = tmp_path / "liste.xlsx"
    lines = [PartLine(("=SUM(A1:A2)", "P"), 1, ("a.dwg",))]

    export_parts(lines, HEADERS, target, None, PartsListSettings())

    cell = load_workbook(target).active
    assert isinstance(cell, Worksheet)
    assert cell["A2"].data_type == "s"
    assert cell["A2"].value == "=SUM(A1:A2)"
