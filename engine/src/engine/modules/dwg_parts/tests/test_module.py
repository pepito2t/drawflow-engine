import json
import subprocess
import sys
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from engine.core.templates import TemplateLibrary, TemplateUser
from engine.parts.tests.plans import build_facade_plan

GOLDEN_RELATIVE_PATH = Path("fixtures") / "dwg-parts" / "expected" / "liste-pieces.json"


def repository_root() -> Path:
    for parent in Path(__file__).resolve().parents:
        if (parent / "fixtures").is_dir() and (parent / "engine").is_dir():
            return parent
    raise RuntimeError("repository root not found")


def run_engine(*arguments: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-m", "engine.cli", *arguments],
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )


def write_json(path: Path, content: dict[str, Any]) -> Path:
    path.write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
    return path


@pytest.fixture
def plans(tmp_path: Path) -> Path:
    folder = tmp_path / "Plans façade"
    folder.mkdir()
    build_facade_plan(folder / "Nord.dxf")
    build_facade_plan(folder / "Sud.dxf")
    return folder


def run_parts_list(tmp_path: Path, plans: Path, batch_size: int) -> list[dict[str, Any]]:
    settings = write_json(tmp_path / "settings.json", {"general": {"batch_size": batch_size}})
    inputs = write_json(
        tmp_path / "inputs.json",
        {
            "folders": [str(plans)],
            "project": "Tour B",
            "output_folder": str(tmp_path / "Sortie"),
            "preview": False,
        },
    )
    completed = run_engine("run", "dwg-parts", "--input", str(inputs), "--settings", str(settings))
    assert completed.returncode == 0, completed.stdout + completed.stderr
    return [json.loads(line) for line in completed.stdout.splitlines()]


@pytest.mark.parametrize("batch_size", [1, 2])
def test_parts_list_matches_golden_file(tmp_path: Path, plans: Path, batch_size: int) -> None:
    events = run_parts_list(tmp_path, plans, batch_size)

    result = events[-1]
    assert result["type"] == "result"
    output = Path(result["outputs"][0])
    assert output.name.startswith("Tour B_liste-pieces_")
    sheet = load_workbook(output).active
    assert isinstance(sheet, Worksheet)
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    golden = json.loads((repository_root() / GOLDEN_RELATIVE_PATH).read_text(encoding="utf-8"))
    assert rows == golden["rows"]


def test_progress_is_reported_per_plan(tmp_path: Path, plans: Path) -> None:
    events = run_parts_list(tmp_path, plans, batch_size=2)

    reads = [event for event in events if event["type"] == "progress"]
    assert [event["current"] for event in reads] == [1, 2]


def test_dwg_without_oda_explains_how_to_configure_it(tmp_path: Path) -> None:
    (tmp_path / "plan.dwg").write_bytes(b"DWG")
    inputs = write_json(
        tmp_path / "inputs.json",
        {"files": [str(tmp_path / "plan.dwg")], "output_folder": str(tmp_path / "Sortie")},
    )

    completed = run_engine("run", "dwg-parts", "--input", str(inputs))

    error = json.loads(completed.stdout.splitlines()[-1])
    assert error["type"] == "error"
    assert "Paramètres" in error["hint"]


def test_inputs_require_a_plan_or_a_folder(tmp_path: Path) -> None:
    inputs = write_json(tmp_path / "inputs.json", {"output_folder": str(tmp_path)})

    completed = run_engine("run", "dwg-parts", "--input", str(inputs))

    error = json.loads(completed.stdout.splitlines()[-1])
    assert "au moins un plan" in error["message"]


def test_default_template_is_used_when_none_is_chosen(tmp_path: Path, plans: Path) -> None:
    template = tmp_path / "Modèle maison.xlsx"
    workbook = Workbook()
    workbook.save(template)
    settings_file = write_json(tmp_path / "settings.json", {})
    library = TemplateLibrary(settings_file)
    imported = library.import_file(template)
    library.set_default(TemplateUser("dwg-parts", "Liste de pièces", "xlsx"), imported.id)
    inputs = write_json(
        tmp_path / "inputs.json",
        {"folders": [str(plans)], "output_folder": str(tmp_path / "Sortie")},
    )

    completed = run_engine(
        "run", "dwg-parts", "--input", str(inputs), "--settings", str(settings_file)
    )

    events = [json.loads(line) for line in completed.stdout.splitlines()]
    assert {"type": "log", "message": "Modèle par défaut : Modèle maison.xlsx"} in events
    assert events[-1]["type"] == "result"


def test_preview_shows_the_table_and_writes_nothing(tmp_path: Path, plans: Path) -> None:
    output = tmp_path / "Sortie"
    inputs = write_json(
        tmp_path / "inputs.json",
        {"folders": [str(plans)], "project": "Test", "output_folder": str(output), "preview": True},
    )

    completed = run_engine("run", "dwg-parts", "--input", str(inputs))
    events = [json.loads(line) for line in completed.stdout.splitlines()]

    [table] = [event for event in events if event["type"] == "table"]
    assert table["headers"][-2:] == ["Quantité", "Plans"]
    assert table["total"] == len(table["rows"]) > 0
    assert events[-1]["type"] == "result" and events[-1]["summary"].startswith("Aperçu")
    assert not output.exists() or not any(output.iterdir())
