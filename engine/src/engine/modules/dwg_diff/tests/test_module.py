import json
import subprocess
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from engine.parts.tests.plans import build_facade_plan, build_facade_plan_revised


def run_engine(*arguments: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-m", "engine.cli", *arguments],
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )


def run_diff(tmp_path: Path, *, preview: bool) -> list[dict[str, Any]]:
    before = tmp_path / "A"
    after = tmp_path / "B"
    before.mkdir()
    after.mkdir()
    build_facade_plan(before / "01_facade.dxf")
    build_facade_plan_revised(after / "02_facade.dxf")
    inputs = tmp_path / "inputs.json"
    inputs.write_text(
        json.dumps(
            {
                "before_folders": [str(before)],
                "after_folders": [str(after)],
                "project": "Tour B",
                "output_folder": str(tmp_path / "Sortie"),
                "preview": preview,
            }
        ),
        encoding="utf-8",
    )
    completed = run_engine("run", "dwg-diff", "--input", str(inputs))
    assert completed.returncode == 0, completed.stdout + completed.stderr
    return [json.loads(line) for line in completed.stdout.splitlines()]


def test_export_lists_every_category_in_its_own_sheet(tmp_path: Path) -> None:
    events = run_diff(tmp_path, preview=False)

    result = events[-1]
    assert result["type"] == "result", events
    # P-1500 added; P-900 and the dynamic window F-80 gone; P-1200 and EQ-40 counts up.
    assert result["summary"] == "1 ajout(s), 2 suppression(s), 2 modification(s), 1 inchangée(s)"
    workbook = load_workbook(result["outputs"][0], read_only=True)
    assert workbook.sheetnames == ["Ajouts", "Suppressions", "Modifications", "Inchangés", "Plans"]
    added = list(workbook["Ajouts"].iter_rows(values_only=True))
    assert added[1][:2] == ("PANNEAU", "P-1500")
    plans = list(workbook["Plans"].iter_rows(values_only=True))
    assert plans[1] == ("facade", "oui", "oui")


def test_preview_shows_the_comparison_without_writing(tmp_path: Path) -> None:
    events = run_diff(tmp_path, preview=True)

    [table] = [event for event in events if event["type"] == "table"]
    assert table["headers"][0] == "Statut"
    assert [row["cells"][0] for row in table["rows"]] == [
        "Ajouté",
        "Supprimé",
        "Supprimé",
        "Modifié",
        "Modifié",
        "Inchangé",
    ]
    assert events[-1]["summary"].startswith("Aperçu")
    assert not (tmp_path / "Sortie").exists()


def test_both_indices_are_required(tmp_path: Path) -> None:
    inputs = tmp_path / "inputs.json"
    inputs.write_text(
        json.dumps({"before_folders": [str(tmp_path)], "output_folder": str(tmp_path)}),
        encoding="utf-8",
    )

    completed = run_engine("run", "dwg-diff", "--input", str(inputs))

    assert completed.returncode == 1
    assert "nouvel indice (B)" in completed.stdout
