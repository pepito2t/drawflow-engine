import json
from pathlib import Path

import pytest
from docx import Document

from engine.assistant.inspect import inspect_file
from engine.core.cache import FileCache, file_digest
from engine.core.errors import EngineError, InvalidInputError
from engine.modules.soumission.tests.workbooks import submission_bytes, truncated_sheet_bytes
from engine.parts.tests.plans import build_facade_plan
from engine.parts.worker import DXF_CACHE_NAMESPACE
from engine.testing.pdf import PdfSpec, TextItem, write_pdf


def _settings_with_cache(tmp_path: Path) -> Path:
    settings = tmp_path / "settings.json"
    payload = {"general": {"cache_folder": str(tmp_path / "cache")}}
    settings.write_text(json.dumps(payload), encoding="utf-8")
    return settings


def test_plan_inspection_counts_blocks_and_lists_attributes(tmp_path: Path) -> None:
    plan = build_facade_plan(tmp_path / "façade.dxf")

    report = inspect_file(tmp_path / "settings.json", str(plan))

    assert report["type"] == "dxf" and report["suggested_feature"] == "dwg-parts"
    assert report["plan_read"] is True
    assert {block["name"] for block in report["blocks"]} >= {"PANNEAU", "EQUERRE"}
    assert "REF" in report["attributes"]


def test_a_dwg_not_converted_yet_is_not_converted_by_the_assistant(tmp_path: Path) -> None:
    dwg = tmp_path / "façade.dwg"
    dwg.write_bytes(b"AC1032")

    report = inspect_file(_settings_with_cache(tmp_path), str(dwg))

    assert report["plan_read"] is False and "Liste de pièces" in report["note"]
    assert "blocks" not in report


def test_a_dwg_already_converted_is_read_from_the_cache(tmp_path: Path) -> None:
    settings = _settings_with_cache(tmp_path)
    dwg = tmp_path / "façade.dwg"
    dwg.write_bytes(b"AC1032")
    converted = build_facade_plan(tmp_path / "converted.dxf")
    entry = FileCache(tmp_path / "cache", DXF_CACHE_NAMESPACE).entry_path(file_digest(dwg), ".dxf")
    entry.parent.mkdir(parents=True)
    entry.write_bytes(converted.read_bytes())

    report = inspect_file(settings, str(dwg))

    assert report["plan_read"] is True
    assert {block["name"] for block in report["blocks"]} >= {"PANNEAU", "EQUERRE"}


def test_pdf_inspection_gives_pages_and_first_text(tmp_path: Path) -> None:
    pdf = write_pdf(
        tmp_path / "plan.pdf",
        PdfSpec(pages=[[TextItem(50, 50, "FACADE NORD"), TextItem(50, 80, "Ind. B")], []]),
    )

    report = inspect_file(tmp_path / "settings.json", str(pdf))

    assert report["pages"] == 2
    assert "FACADE NORD" in report["first_page_text"]
    assert report["scanned"] is False and report["suggested_feature"] == "pdf-report"


def test_workbook_and_document_inspection(tmp_path: Path) -> None:
    workbook = tmp_path / "offre.xlsx"
    workbook.write_bytes(submission_bytes())
    document = Document()
    document.add_paragraph("Rapport de chantier")
    document.add_paragraph("Façade nord, indice B")
    docx = tmp_path / "rapport.docx"
    document.save(str(docx))

    sheets = inspect_file(tmp_path / "settings.json", str(workbook))["sheets"]
    paragraphs = inspect_file(tmp_path / "settings.json", str(docx))["paragraphs"]

    assert sheets[0]["sheet"] == "Offre"
    assert any("Désignation" in row for row in sheets[0]["first_rows"])
    assert paragraphs == ["Rapport de chantier", "Façade nord, indice B"]


def test_unknown_or_missing_files_are_refused(tmp_path: Path) -> None:
    notes = tmp_path / "notes.txt"
    notes.write_text("x", encoding="utf-8")

    with pytest.raises(InvalidInputError, match="non pris en charge"):
        inspect_file(tmp_path / "settings.json", str(notes))
    with pytest.raises(InvalidInputError, match="introuvable"):
        inspect_file(tmp_path / "settings.json", str(tmp_path / "absent.dwg"))


def test_damaged_workbook_and_document_are_readable_errors(tmp_path: Path) -> None:
    workbook = tmp_path / "tronqué.xlsx"
    workbook.write_bytes(truncated_sheet_bytes())
    document = tmp_path / "cassé.docx"
    document.write_text("pas un document", encoding="utf-8")

    with pytest.raises(EngineError) as workbook_error:
        inspect_file(tmp_path / "settings.json", str(workbook))
    with pytest.raises(InvalidInputError) as document_error:
        inspect_file(tmp_path / "settings.json", str(document))

    assert workbook_error.value.file == workbook and "« Offre »" in workbook_error.value.message
    assert document_error.value.file == document and document_error.value.hint is not None
