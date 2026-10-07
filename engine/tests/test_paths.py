import errno
import sys
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
from docxtpl import DocxTemplate
from openpyxl import Workbook, load_workbook

from engine.core import paths
from engine.core.errors import OutputWriteError
from engine.core.events import ResultEvent
from engine.core.json_files import replace_file
from engine.core.naming import (
    MAX_STEM_LENGTH,
    MIN_STEM_LENGTH,
    PARTIAL_SUFFIX,
    UNIQUE_INDEX_RESERVE,
    OutputFolderError,
    output_target,
    stem_length_for,
    unique_output_path,
    writing_output,
)
from engine.core.paths import (
    EXTENDED_PREFIX,
    WINDOWS_MAX_PATH_LENGTH,
    extended_form,
    extended_path,
    write_failure_text,
)
from engine.core.registry import get_module
from engine.core.runner import run_module
from engine.core.settings import load_run_settings
from engine.core.xlsx import save_workbook
from engine.modules.pdf_report.docx_render import render_report
from engine.modules.pdf_report.report import build_context
from engine.modules.pdf_report.settings import DEFAULT_FIELDS
from engine.modules.soumission.tests.workbooks import write_submission

UNC_FOLDER = r"\\serveur\partage\Façades 2026"
DEEP_SEGMENT = "Dossier très profond é"
DEEP_SEGMENT_COUNT = 14
TOO_LONG_HINT = "caractères"
MOMENT = datetime(2026, 10, 7, 9, 0)

type Recorded = list[str]


def windows_deep_path() -> str:
    return "C:\\" + "\\".join([DEEP_SEGMENT] * DEEP_SEGMENT_COUNT) + "\\liste.xlsx"


def deep_folder(base: Path) -> Path:
    folder = base.joinpath(*[DEEP_SEGMENT] * DEEP_SEGMENT_COUNT)
    assert len(str(folder)) > WINDOWS_MAX_PATH_LENGTH
    return folder


@pytest.fixture
def on_windows(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(sys, "platform", "win32")


def record_path_calls(monkeypatch: pytest.MonkeyPatch, name: str) -> Recorded:
    """Records every path a `Path` method touches, without touching the disk."""
    calls: Recorded = []

    def fake(self: Path, *arguments: Any, **_: Any) -> None:
        calls.extend(str(path) for path in (self, *arguments) if isinstance(path, Path))

    monkeypatch.setattr(Path, name, fake)
    return calls


def record_saves(monkeypatch: pytest.MonkeyPatch, owner: type) -> Recorded:
    calls: Recorded = []

    def fake(_: object, destination: object) -> None:
        calls.append(str(destination))

    monkeypatch.setattr(owner, "save", fake)
    return calls


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (
            r"C:\Users\Zoé Dupont\Offres 2026\liste.xlsx",
            r"\\?\C:\Users\Zoé Dupont\Offres 2026\liste.xlsx",
        ),
        (UNC_FOLDER + r"\liste.xlsx", r"\\?\UNC\serveur\partage\Façades 2026\liste.xlsx"),
        ("C:/Projets/../Sortie/liste.xlsx", r"\\?\C:\Sortie\liste.xlsx"),
        (r"\\?\C:\déjà\préfixé.xlsx", r"\\?\C:\déjà\préfixé.xlsx"),
        (r"\\.\pipe\drawflow", r"\\.\pipe\drawflow"),
        (r"Sortie\liste.xlsx", r"Sortie\liste.xlsx"),
        ("C:relatif.xlsx", "C:relatif.xlsx"),
    ],
)
def test_extended_form_prefixes_only_absolute_windows_paths(raw: str, expected: str) -> None:
    assert extended_form(raw) == expected


def test_extended_form_keeps_a_path_longer_than_max_path_whole() -> None:
    raw = windows_deep_path()

    assert len(raw) > WINDOWS_MAX_PATH_LENGTH
    assert extended_form(raw) == EXTENDED_PREFIX + raw


def test_extended_path_is_identity_outside_windows(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(sys, "platform", "linux")
    path = Path(UNC_FOLDER)

    assert extended_path(path) is path


@pytest.mark.usefixtures("on_windows")
def test_extended_path_on_windows_prefixes_unc_shares() -> None:
    assert str(extended_path(Path(UNC_FOLDER))) == r"\\?\UNC\serveur\partage\Façades 2026"


@pytest.mark.usefixtures("on_windows")
def test_output_name_is_reserved_through_the_extended_form_and_returned_plain(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    created = record_path_calls(monkeypatch, "mkdir")
    touched = record_path_calls(monkeypatch, "touch")

    target = unique_output_path(Path(UNC_FOLDER), "liste é.xlsx")

    assert created == [r"\\?\UNC\serveur\partage\Façades 2026"]
    assert touched == [r"\\?\UNC\serveur\partage\Façades 2026\liste é.xlsx"]
    assert not str(target).startswith(EXTENDED_PREFIX)


@pytest.mark.usefixtures("on_windows")
def test_replace_and_saves_use_the_extended_form(monkeypatch: pytest.MonkeyPatch) -> None:
    replaced = record_path_calls(monkeypatch, "replace")
    record_path_calls(monkeypatch, "mkdir")
    saved = record_saves(monkeypatch, Workbook)
    target = Path(windows_deep_path())

    replace_file(Path(windows_deep_path() + PARTIAL_SUFFIX), target)
    save_workbook(Workbook(), target)

    assert replaced == [
        EXTENDED_PREFIX + windows_deep_path() + PARTIAL_SUFFIX,
        EXTENDED_PREFIX + windows_deep_path(),
    ]
    assert saved == [EXTENDED_PREFIX + windows_deep_path()]


@pytest.mark.usefixtures("on_windows")
def test_report_is_saved_through_the_extended_form(monkeypatch: pytest.MonkeyPatch) -> None:
    record_path_calls(monkeypatch, "mkdir")
    saved = record_saves(monkeypatch, DocxTemplate)
    target = Path(UNC_FOLDER + r"\rapport.docx")

    render_report(None, build_context([], DEFAULT_FIELDS, "Tour B", MOMENT), target)

    assert saved == [r"\\?\UNC\serveur\partage\Façades 2026\rapport.docx"]


def test_a_name_too_long_for_the_system_is_named_as_the_cause(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    def refuse(self: Path, *_: Any, **__: Any) -> None:
        raise OSError(errno.ENAMETOOLONG, "File name too long", str(self))

    monkeypatch.setattr(Path, "touch", refuse)

    with pytest.raises(OutputFolderError) as caught:
        unique_output_path(tmp_path, "liste.xlsx")

    assert caught.value.message == "Le chemin du fichier de sortie est trop long."
    assert caught.value.hint is not None and TOO_LONG_HINT in caught.value.hint


def test_windows_filename_range_error_is_named_as_too_long() -> None:
    error = OSError(errno.ENOENT, "introuvable")
    error.winerror = paths.ERROR_FILENAME_EXCED_RANGE  # type: ignore[attr-defined]  # Windows only.

    message, _ = write_failure_text(Path("C:/court.xlsx"), error, "générique", "aide")

    assert message == "Le chemin du fichier de sortie est trop long."


@pytest.mark.usefixtures("on_windows")
def test_a_path_beyond_max_path_on_windows_is_named_as_too_long() -> None:
    error = PermissionError(errno.EACCES, "refusé")

    message, hint = write_failure_text(Path(windows_deep_path()), error, "générique", "aide")

    assert message == "Le chemin du fichier de sortie est trop long."
    assert str(len(windows_deep_path())) in hint


def test_an_ordinary_failure_keeps_its_own_message(tmp_path: Path) -> None:
    error = PermissionError(errno.EACCES, "refusé")

    assert write_failure_text(tmp_path / "a.xlsx", error, "générique", "aide") == (
        "générique",
        "aide",
    )


def folder_of_length(length: int) -> Path:
    anchor = Path.cwd().anchor
    return Path(anchor + "d" * (length - len(anchor)))


@pytest.mark.parametrize(
    ("folder_length", "expected"),
    [(20, MAX_STEM_LENGTH), (100, 139), (230, MIN_STEM_LENGTH)],
)
def test_name_length_follows_the_room_left_by_the_folder(folder_length: int, expected: int) -> None:
    assert stem_length_for(folder_of_length(folder_length), "xlsx") == expected


def test_output_path_with_spaces_and_accents_stays_within_max_path(tmp_path: Path) -> None:
    padding = max(1, 150 - len(str(tmp_path)) - 1)
    folder = tmp_path / ("Sortie é " * padding)[:padding].strip()

    target = output_target(folder, "{projet}", {"projet": "Façade nord " * 40}, "xlsx")

    longest = len(str(target)) + UNIQUE_INDEX_RESERVE + len(PARTIAL_SUFFIX)
    assert longest <= WINDOWS_MAX_PATH_LENGTH


def test_writes_into_a_folder_deeper_than_max_path(tmp_path: Path) -> None:
    folder = deep_folder(tmp_path)

    target = output_target(folder, "{projet}", {"projet": "Tour B é"}, "xlsx")
    with writing_output(target) as draft:
        save_workbook(Workbook(), draft)

    assert not str(target).startswith(EXTENDED_PREFIX)
    assert load_workbook(extended_path(target)).sheetnames


def test_save_failure_in_a_missing_unc_share_is_typed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def refuse(self: Path, *_: Any, **__: Any) -> None:
        raise FileNotFoundError(errno.ENOENT, "partage absent", str(self))

    monkeypatch.setattr(Path, "mkdir", refuse)

    with pytest.raises(OutputWriteError) as caught:
        save_workbook(Workbook(), Path(UNC_FOLDER + r"\liste.xlsx"))

    assert caught.value.file == Path(UNC_FOLDER + r"\liste.xlsx")


def visible_extended_path(link: Path, calls: Recorded) -> Callable[[Path], Path]:
    """Outside Windows: a real, working path that carries the prefix through a symbolic link."""

    def extended(path: Path) -> Path:
        calls.append(str(path))
        return Path(f"{link}{path.absolute()}")

    return extended


def use_visible_extended_paths(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Recorded:
    calls: Recorded = []
    if sys.platform == "win32":
        return calls
    link = tmp_path / EXTENDED_PREFIX
    link.symlink_to("/", target_is_directory=True)
    fake = visible_extended_path(link, calls)
    original = paths.extended_path
    for module in list(sys.modules.values()):
        if getattr(module, "extended_path", None) is original:
            monkeypatch.setattr(module, "extended_path", fake)
    return calls


def test_result_paths_never_carry_the_extended_prefix(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    calls = use_visible_extended_paths(monkeypatch, tmp_path)
    source = write_submission(tmp_path / "Offre façadier é.xlsx")
    module = get_module("soumission")
    inputs = {"files": [str(source)], "output_folder": str(deep_folder(tmp_path)), "preview": False}
    events: list[object] = []

    run_module(module, inputs, load_run_settings(None, module), events.append)

    results = [event for event in events if isinstance(event, ResultEvent)]
    outputs = [output for result in results for output in result.outputs]
    assert outputs and all(EXTENDED_PREFIX not in output for output in outputs)
    assert all(EXTENDED_PREFIX not in str(event) for event in events)
    assert extended_path(Path(outputs[0])).is_file()
    assert sys.platform == "win32" or calls
